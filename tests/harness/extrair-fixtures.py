#!/usr/bin/env python3
"""
tests/harness/extrair-fixtures.py — extrai as abas que
montarSerieHistoricoInicio_/calcularFluxoCaixaDiario_ (apps-script/) leem
de um export .xlsx real da planilha, no MESMO formato que
SpreadsheetApp.getRange(...).getValues() devolveria (array de arrays,
1 linha por row, datas viram Date de verdade no harness) — pra
tests/harness/gas-vm-harness.mjs rodar o código de produção sem mudar
nada nele. Ver comentário no topo do harness.

Uso:
  python3 tests/harness/extrair-fixtures.py "/caminho/Investimentos - Controle NN.xlsx" [saida.json]

Precisa de openpyxl (`pip install openpyxl --break-system-packages` se
não tiver).
"""
import sys
import os
import json
import datetime

try:
    import openpyxl
except ImportError:
    print('openpyxl não instalado. Rode: pip install openpyxl --break-system-packages', file=sys.stderr)
    sys.exit(1)

SHEETS = [
    'aux_historico-patrimonio',
    'aux_historico-renda-fixa',
    'aux_historico-indices',
    'Transações',
    'Transações - USA',
    'Transações Renda Fixa',
    'Carteira Renda Fixa',
    'Proventos',
    'Proventos - USA',
    # --- 20/09/2026: acrescentadas pra rodar as funções LIVE das telas de
    # Carteiras (montarHome_/montarCarteirasHome_/montarCarteiraClasse_/
    # montarCarteirasRendaFixa_) contra dados reais, a pedido do Tiago
    # ("veja se os números lá batem também") ---
    '📊Dash Geral',
    'Auxiliar_app',
    'Distribuição e Metas',
    'Auxiliar_ativos',
    'Carteira FIIs',
    'RF Contratada - Resumo',
    'RF Contratada - Lotes',
    # --- 26/09/2026: tela Organização Financeira (Despesas.gs) ---
    'Despesas Essenciais',
]


def serialize(v):
    if isinstance(v, datetime.datetime):
        return {'__date__': v.isoformat()}
    if isinstance(v, datetime.date):
        return {'__date__': v.isoformat() + 'T00:00:00'}
    if isinstance(v, datetime.time):
        return {'__date__': '1899-12-30T' + v.isoformat()}
    if v is None:
        return None
    return v


def main():
    if len(sys.argv) < 2:
        print('uso: extrair-fixtures.py <planilha.xlsx> [saida.json]', file=sys.stderr)
        sys.exit(1)
    xlsx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), 'fixtures.json')

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    out = {}
    for nome in SHEETS:
        if nome not in wb.sheetnames:
            print(f'AVISO: aba "{nome}" não encontrada na planilha - pulando (fixture ficará ausente).', file=sys.stderr)
            continue
        ws = wb[nome]
        linhas = [[serialize(c) for c in row] for row in ws.iter_rows(values_only=True)]
        last_row_idx = 0
        for i, linha in enumerate(linhas, start=1):
            if any(c is not None for c in linha):
                last_row_idx = i
        out[nome] = {'linhas': linhas[:last_row_idx], 'lastRow': last_row_idx}
        print(nome, '-> lastRow', last_row_idx, 'cols', len(linhas[0]) if linhas else 0)

    # 23/09/2026 #4: de onde as fixtures vieram - o relatório de conferência
    # das telas (relatorio-telas.mjs) mostra isso no topo. A "data dos
    # dados" de verdade (último sync de preços) é lida das próprias linhas.
    # Chaves com "_" na frente não são abas (o harness pula).
    out['_meta'] = {
        'arquivo': os.path.basename(xlsx_path),
        'arquivoModificadoEm': datetime.datetime.fromtimestamp(os.path.getmtime(xlsx_path), datetime.timezone.utc).isoformat(),
        'extraidoEm': datetime.datetime.now(datetime.timezone.utc).isoformat(),
    }

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False)

    print('gravado', out_path, '-', os.path.getsize(out_path), 'bytes')


if __name__ == '__main__':
    main()
